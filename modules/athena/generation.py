"""
modules/athena/generation.py

Report generation for Athena (backlog #53, #51, #52, #66).

One report model, several renderers
-----------------------------------
``ReportModel`` (title, sections, references) is built once, from the
literature review Athena can already write plus the file-based citations
the citation manager produces. Every output format renders that same
model, so a fix to content shows up in all of them:

* ``render_pdf``    -- reportlab                    (#53)
* ``render_latex``  -- a ``.tex`` plus a ``.bib``    (#51)
* ``render_pptx``   -- python-pptx slide deck        (#52)

``generate_methodology`` (#66) asks the model for a JSON research-method
skeleton, validates it strictly, and turns it into a ``ReportModel`` so it
goes through the same renderers.

Honest limits
-------------
* References are FILE-based (see citation_service.py): no author or year
  is invented.
* When the model is down or returns junk, nothing half-written is saved:
  callers get ``None`` / a ``GenerationError`` carrying a plain message.
* The LaTeX output is a scaffold meant to be edited, not a finished paper.
"""
from __future__ import annotations

import json
import logging
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Callable, Optional

logger = logging.getLogger(__name__)

FORMATS = ("pdf", "latex", "pptx")
_FORMAT_ALIASES = {
    "pdf": "pdf",
    "latex": "latex", "tex": "latex",
    "pptx": "pptx", "ppt": "pptx", "powerpoint": "pptx", "slides": "pptx",
    "slide": "pptx", "presentation": "pptx", "deck": "pptx",
}
_EXTENSIONS = {"pdf": ".pdf", "latex": ".tex", "pptx": ".pptx"}

_MAX_BULLETS = 5
_MAX_BULLET_CHARS = 180


class GenerationError(Exception):
    """A failure with a message that is safe to say to the user."""


# ---------------------------------------------------------------------------
# Model
# ---------------------------------------------------------------------------

@dataclass
class Section:
    heading: str
    body: str                       # paragraphs separated by blank lines


@dataclass
class Reference:
    key: str                        # BibTeX key
    title: str
    file_name: str
    pages: str = ""                 # e.g. "3-7", or ""


@dataclass
class ReportModel:
    title: str
    sections: list[Section] = field(default_factory=list)
    references: list[Reference] = field(default_factory=list)
    subtitle: str = ""
    generated_on: str = field(default_factory=lambda: date.today().isoformat())

    def paragraphs(self, section: Section) -> list[str]:
        return [p.strip() for p in re.split(r"\n\s*\n", section.body) if p.strip()]


def normalise_format(value: Optional[str], default: str = "pdf") -> str:
    """Canonical format name; raises GenerationError for one we can't make."""
    if not value or not str(value).strip():
        return default
    key = str(value).strip().lower().lstrip(".")
    if key not in _FORMAT_ALIASES:
        raise GenerationError(
            f"I can't make a \"{value}\" file. I can make a PDF, a LaTeX file or a PowerPoint."
        )
    return _FORMAT_ALIASES[key]


def build_references(sources: list[dict[str, Any]]) -> list[Reference]:
    """File-based references from search-source dicts or {"file_name"} dicts."""
    from modules.athena.services.citation_service import CitationRegistry

    registry = CitationRegistry()
    registry.add_sources(sources)
    refs: list[Reference] = []
    for i, c in enumerate(registry.citations(), start=1):
        stem = "".join(ch for ch in Path(c.file_name).stem if ch.isalnum())[:20] or "file"
        pages = f"{min(c.pages)}-{max(c.pages)}" if c.pages else ""
        refs.append(Reference(key=f"src{i}_{stem}", title=c.title_guess,
                              file_name=c.file_name, pages=pages))
    return refs


# ---------------------------------------------------------------------------
# Building a report from Athena's material
# ---------------------------------------------------------------------------

_FAILURE_PREFIXES = ("i don't have any documents", "i had trouble", "i couldn't")


def review_failed(review: str) -> bool:
    """SynthesisService reports failure as a plain sentence; detect it."""
    low = (review or "").strip().lower()
    return not low or low.startswith(_FAILURE_PREFIXES)


def build_report_from_subject(synthesis, subject: str, title: Optional[str] = None) -> ReportModel:
    """
    Report on *subject* from the literature review plus the reference list.
    Raises GenerationError (with a user-safe message) when there is nothing
    good to build from; never returns an empty report.
    """
    result = synthesis.generate_literature_review(subject)
    review = (result.get("review") or "").strip()
    source_names = result.get("sources") or []
    if not source_names:
        raise GenerationError(
            f"I don't have any documents indexed under \"{subject}\", so there is nothing to report on."
        )
    if review_failed(review):
        raise GenerationError(
            f"I couldn't write the review for \"{subject}\" just now, so I didn't create a file. "
            "Please try again in a moment."
        )
    report = ReportModel(
        title=title or f"Literature review: {subject}",
        subtitle=f"Based on {len(source_names)} indexed document(s)",
        sections=[Section("Overview", review)],
        references=build_references([{"file_name": n, "subject": subject} for n in source_names]),
    )
    return report


# ---------------------------------------------------------------------------
# Methodology (#66)
# ---------------------------------------------------------------------------

_METHOD_FIELDS = ("research_question", "design", "independent_variables",
                  "dependent_variables", "controls", "data_collection", "analysis")
_LIST_FIELDS = {"independent_variables", "dependent_variables", "controls", "hypotheses", "limitations"}

_METHOD_PROMPT = (
    "You are helping design a research methodology. Reply with ONLY a JSON object, "
    "no prose and no code fences, with exactly these keys: "
    '"research_question" (string), "hypotheses" (list of strings), '
    '"design" (string), "independent_variables" (list of strings), '
    '"dependent_variables" (list of strings), "controls" (list of strings), '
    '"data_collection" (string), "analysis" (string), "limitations" (list of strings).\n\n'
    "Topic / question: {question}"
)


def _extract_json_object(text: str) -> Optional[dict]:
    """The first JSON object in *text*, tolerating code fences and chatter."""
    if not text:
        return None
    cleaned = re.sub(r"```(?:json)?", "", text)
    start = cleaned.find("{")
    while start != -1:
        depth = 0
        in_str = False
        esc = False
        for i in range(start, len(cleaned)):
            ch = cleaned[i]
            if in_str:
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == '"':
                    in_str = False
                continue
            if ch == '"':
                in_str = True
            elif ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    try:
                        obj = json.loads(cleaned[start:i + 1])
                        return obj if isinstance(obj, dict) else None
                    except ValueError:
                        break
        start = cleaned.find("{", start + 1)
    return None


def validate_methodology(obj: Any) -> Optional[dict[str, Any]]:
    """
    Normalised methodology dict, or None if it isn't usable. Required: a
    research question, a design, and at least one independent and one
    dependent variable. Lists are coerced from a single string if needed.
    """
    if not isinstance(obj, dict):
        return None
    out: dict[str, Any] = {}
    for key in _METHOD_FIELDS + ("hypotheses", "limitations"):
        val = obj.get(key)
        if key in _LIST_FIELDS:
            if isinstance(val, str):
                val = [val]
            if not isinstance(val, list):
                val = []
            out[key] = [str(v).strip() for v in val if str(v).strip()]
        else:
            out[key] = str(val).strip() if isinstance(val, (str, int, float)) else ""
    required_text = ("research_question", "design")
    required_lists = ("independent_variables", "dependent_variables")
    if any(not out[k] for k in required_text) or any(not out[k] for k in required_lists):
        return None
    return out


def methodology_to_report(m: dict[str, Any]) -> ReportModel:
    def bullets(items: list[str]) -> str:
        return "\n\n".join(f"- {i}" for i in items)

    sections = [
        Section("Research question", m["research_question"]),
    ]
    if m.get("hypotheses"):
        sections.append(Section("Hypotheses", bullets(m["hypotheses"])))
    sections.append(Section("Study design", m["design"]))
    sections.append(Section("Variables", "Independent:\n\n" + bullets(m["independent_variables"])
                            + "\n\nDependent:\n\n" + bullets(m["dependent_variables"])))
    if m.get("controls"):
        sections.append(Section("Controls", bullets(m["controls"])))
    if m.get("data_collection"):
        sections.append(Section("Data collection", m["data_collection"]))
    if m.get("analysis"):
        sections.append(Section("Analysis plan", m["analysis"]))
    if m.get("limitations"):
        sections.append(Section("Limitations", bullets(m["limitations"])))
    return ReportModel(title="Research methodology", subtitle=m["research_question"][:120],
                       sections=sections)


def generate_methodology(ai, question: str) -> dict[str, Any]:
    """
    Ask the model for a methodology skeleton and validate it. *ai* is a
    HestiaLLMAdapter-like object (``generate(prompt) -> {"text","error"}``).
    Raises GenerationError with a clear message if the model is down or
    returns something that isn't a usable skeleton.
    """
    question = (question or "").strip()
    if not question:
        raise GenerationError("What topic or research question should the methodology be for?")
    try:
        res = ai.generate(_METHOD_PROMPT.format(question=question))
    except Exception:
        logger.exception("Methodology generation: LLM call failed.")
        res = {"text": "", "error": "exception"}
    text = res.get("text", "") if isinstance(res, dict) else str(res)
    error = res.get("error") if isinstance(res, dict) else None
    if error or not (text or "").strip():
        raise GenerationError("I couldn't reach the language model to draft that methodology. "
                              "Please try again in a moment.")
    valid = validate_methodology(_extract_json_object(text))
    if valid is None:
        raise GenerationError("The model's answer wasn't a usable methodology, so I didn't create "
                              "a file. Trying again, or rewording the question, usually helps.")
    return valid


# ---------------------------------------------------------------------------
# Slide outline (#52)
# ---------------------------------------------------------------------------

_SENTENCE_RE = re.compile(r"(?<=[.!?])\s+")


def _bulletize(body: str) -> list[str]:
    """Deterministic bullets: list items as written, otherwise first sentences."""
    items: list[str] = []
    for para in re.split(r"\n\s*\n", body):
        para = para.strip()
        if not para:
            continue
        if para.startswith(("- ", "* ")):
            items.append(para[2:].strip())
        else:
            items.extend(s.strip() for s in _SENTENCE_RE.split(para) if s.strip())
    out = []
    for it in items:
        if len(it) > _MAX_BULLET_CHARS:
            it = it[:_MAX_BULLET_CHARS].rsplit(" ", 1)[0].rstrip(",;:") + "\u2026"
        out.append(it)
    return out[:_MAX_BULLETS]


def deterministic_outline(model: ReportModel) -> list[dict[str, Any]]:
    return [{"title": s.heading, "bullets": _bulletize(s.body)} for s in model.sections]


def outline_with_llm(ai, model: ReportModel) -> list[dict[str, Any]]:
    """
    Ask the model for tighter slide bullets; fall back to the deterministic
    outline on ANY problem (model down, junk JSON, empty result).
    """
    body = "\n\n".join(f"{s.heading}\n{s.body}" for s in model.sections)[:6000]
    prompt = (
        "Turn this report into a slide outline. Reply with ONLY a JSON object: "
        '{"slides": [{"title": str, "bullets": [str, ...]}]}. '
        f"At most {_MAX_BULLETS} short bullets per slide.\n\n{body}"
    )
    try:
        res = ai.generate(prompt)
        text = res.get("text", "") if isinstance(res, dict) else str(res)
        obj = _extract_json_object(text)
        slides = obj.get("slides") if obj else None
        cleaned = []
        for s in slides or []:
            if not isinstance(s, dict):
                continue
            title = str(s.get("title") or "").strip()
            bl = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()]
            if title and bl:
                cleaned.append({"title": title, "bullets": bl[:_MAX_BULLETS]})
        if cleaned:
            return cleaned
    except Exception:
        logger.debug("LLM slide outline failed; using the deterministic one.", exc_info=True)
    return deterministic_outline(model)


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------

def _ref_line(r: Reference) -> str:
    pages = f" (pp. {r.pages})" if r.pages else ""
    return f"{r.title}{pages}. [{r.file_name}]"


def _xml_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_pdf(model: ReportModel, path: str) -> str:
    """Write *model* to a PDF at *path* (reportlab). Returns the path."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph(_xml_escape(model.title), styles["Title"])]
    if model.subtitle:
        story.append(Paragraph(_xml_escape(model.subtitle), styles["Italic"]))
    story.append(Paragraph(model.generated_on, styles["Normal"]))
    story.append(Spacer(1, 18))
    for sec in model.sections:
        story.append(Paragraph(_xml_escape(sec.heading), styles["Heading2"]))
        for para in model.paragraphs(sec):
            if para.startswith(("- ", "* ")):
                story.append(ListFlowable(
                    [ListItem(Paragraph(_xml_escape(para[2:]), styles["Normal"]))],
                    bulletType="bullet"))
            else:
                story.append(Paragraph(_xml_escape(para).replace("\n", "<br/>"), styles["Normal"]))
            story.append(Spacer(1, 6))
    if model.references:
        story.append(Paragraph("References", styles["Heading2"]))
        for r in model.references:
            story.append(Paragraph(_xml_escape(_ref_line(r)), styles["Normal"]))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, title=model.title).build(story)
    return str(path)


_LATEX_SPECIALS = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
    "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
}


def latex_escape(text: str) -> str:
    return "".join(_LATEX_SPECIALS.get(ch, ch) for ch in text)


def _bib_value(text: str) -> str:
    """A value that is safe inside BibTeX braces."""
    return latex_escape(text)


def render_latex(model: ReportModel, tex_path: str) -> tuple[str, str]:
    """
    Write ``<name>.tex`` and, when there are references, ``<name>.bib``
    next to it. Returns (tex_path, bib_path); bib_path is "" if no .bib was
    needed. The .tex uses plain ``article`` and ``\\nocite{*}``
    so every reference is listed even before you add \\cite commands.
    """
    tex = Path(tex_path)
    bib = tex.with_suffix(".bib")
    tex.parent.mkdir(parents=True, exist_ok=True)

    lines = [
        r"\documentclass[11pt]{article}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage[margin=2.5cm]{geometry}",
        r"\usepackage{enumitem}",
        "",
        rf"\title{{{latex_escape(model.title)}}}",
        rf"\author{{{latex_escape(model.subtitle) if model.subtitle else ''}}}",
        rf"\date{{{model.generated_on}}}",
        "",
        r"\begin{document}",
        r"\maketitle",
        "",
    ]
    for sec in model.sections:
        lines.append(rf"\section{{{latex_escape(sec.heading)}}}")
        items: list[str] = []

        def flush() -> None:
            if items:
                lines.append(r"\begin{itemize}[noitemsep]")
                lines.extend(rf"  \item {i}" for i in items)
                lines.append(r"\end{itemize}")
                items.clear()

        for para in model.paragraphs(sec):
            if para.startswith(("- ", "* ")):
                items.append(latex_escape(para[2:]))
            else:
                flush()
                lines.append(latex_escape(para).replace("\n", " "))
            lines.append("")
        flush()
    if model.references:
        lines += [r"\nocite{*}", r"\bibliographystyle{plain}", rf"\bibliography{{{bib.stem}}}"]
    lines += ["", r"\end{document}", ""]
    tex.write_text("\n".join(lines), encoding="utf-8")

    if not model.references:
        return str(tex), ""
    entries = []
    for r in model.references:
        note = f"Source file: {r.file_name}" + (f", pages {r.pages}" if r.pages else "")
        entries.append(f"@misc{{{r.key},\n  title = {{{_bib_value(r.title)}}},\n"
                       f"  note = {{{_bib_value(note)}}}\n}}")
    bib.write_text("\n\n".join(entries) + "\n", encoding="utf-8")
    return str(tex), str(bib)


def render_pptx(model: ReportModel, path: str,
                outline: Optional[list[dict[str, Any]]] = None) -> str:
    """Write a slide deck: title slide, one slide per outline entry, sources slide."""
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = model.title
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = model.subtitle or model.generated_on

    for entry in (outline if outline is not None else deterministic_outline(model)):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = entry["title"]
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for i, bullet in enumerate(entry["bullets"] or [""]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = bullet
            para.font.size = Pt(20)

    if model.references:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Sources"
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for i, r in enumerate(model.references[:8]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = _ref_line(r)
            para.font.size = Pt(14)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return str(path)


# ---------------------------------------------------------------------------
# One entry point for the engine
# ---------------------------------------------------------------------------

def slugify(text: str, max_len: int = 40) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")
    return (slug[:max_len].strip("-")) or "report"


def render(model: ReportModel, fmt: str, out_dir: str, stem: str,
           outline_fn: Optional[Callable[[ReportModel], list[dict[str, Any]]]] = None) -> list[str]:
    """Render *model* as *fmt* into *out_dir*; returns the file paths created."""
    fmt = normalise_format(fmt)
    base = Path(out_dir) / f"{slugify(stem)}-{model.generated_on}"
    try:
        if fmt == "pdf":
            return [render_pdf(model, str(base) + ".pdf")]
        if fmt == "latex":
            return [p for p in render_latex(model, str(base) + ".tex") if p]
        return [render_pptx(model, str(base) + ".pptx", outline_fn(model) if outline_fn else None)]
    except GenerationError:
        raise
    except ImportError as exc:
        raise GenerationError(
            f"I can't make that file because the {exc.name or 'required'} library isn't installed."
        ) from exc
    except Exception as exc:
        logger.exception("Rendering %s failed", fmt)
        raise GenerationError("I hit a problem writing that file, so nothing was saved.") from exc


# ---------------------------------------------------------------------------
# Raw-query fallback (alias-routed requests arrive with no entities)
# ---------------------------------------------------------------------------

_FORMAT_WORDS = (r"pdf|latex|tex|powerpoint|pptx|ppt|slides|slide deck|presentation|deck")
_FORMAT_IN_QUERY = re.compile(rf"\b({_FORMAT_WORDS})\b", re.I)
_LEAD_IN = re.compile(
    r"^\s*(?:please\s+)?(?:can you\s+)?(?:generate|make|create|write|build|export|turn|draft|design)\s+"
    r"(?:me\s+)?(?:an?\s+|the\s+|my\s+)?",
    re.I,
)


def parse_report_request(raw: str) -> tuple[str, Optional[str]]:
    """
    (subject, format) from phrasings like "make a pdf report on heat transfer"
    or "turn my machine learning notes into a powerpoint". Either may be
    empty/None when it can't be found.
    """
    text = (raw or "").strip().rstrip("?.! ")
    fmt_match = _FORMAT_IN_QUERY.search(text)
    fmt = fmt_match.group(1).lower() if fmt_match else None
    if fmt == "slide deck":
        fmt = "slides"

    subject = ""
    m = re.search(r"\b(?:on|about|covering)\s+(.+)$", text, re.I)
    if m:
        subject = m.group(1)
    else:
        m = re.search(r"\bturn\s+(?:my\s+|the\s+)?(.+?)(?:\s+notes|\s+documents|\s+papers)?\s+into\b", text, re.I)
        if m:
            subject = m.group(1)
    # "... on heat transfer as a pdf" / "in latex": drop a trailing format phrase.
    subject = re.sub(rf"\s+(?:as|in|into|to)\s+(?:an?\s+)?(?:{_FORMAT_WORDS})\b.*$", "", subject, flags=re.I)
    subject = re.sub(rf"^(?:{_FORMAT_WORDS})\s+(?:report\s+)?(?:on|about)\s+", "", subject, flags=re.I)
    return subject.strip(" ,"), fmt


def parse_methodology_request(raw: str) -> str:
    """The topic after "methodology for/on/about", or the query minus its lead-in."""
    text = (raw or "").strip().rstrip("?.! ")
    m = re.search(r"\bmethodology\s+(?:for|on|about|to)\s+(.+)$", text, re.I)
    if m:
        topic = m.group(1)
    else:
        topic = _LEAD_IN.sub("", text)
        topic = re.sub(r"^methodology\s*", "", topic, flags=re.I)
    topic = re.sub(rf"\s+(?:as|in|into|to)\s+(?:an?\s+)?(?:{_FORMAT_WORDS})\b.*$", "", topic, flags=re.I)
    return topic.strip(" ,")
