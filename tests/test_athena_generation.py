# tests/test_athena_generation.py
"""
Tests for modules/athena/generation.py (backlog #53 PDF, #51 LaTeX,
#52 PowerPoint, #66 methodology) and the two engine intents that use it.

Every renderer is checked by opening its own output with a real reader
(pdfplumber, python-pptx, pdflatex). The model is faked; no network.
"""
import os
import shutil
import subprocess
import sys
import tempfile

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.dirname(__file__))

from test_athena import make_engine  # noqa: E402
from modules.athena import generation as gen  # noqa: E402


def _model():
    return gen.ReportModel(
        title="Literature review: heat transfer & flow",
        subtitle="Based on 2 indexed document(s)",
        sections=[
            gen.Section("Overview", "Conduction is 100% driven by gradients. Convection adds flow. "
                                    "Radiation needs no medium.\n\nSecond paragraph with C# and $5 and _under_."),
            gen.Section("Variables", "Independent:\n\n- temperature\n\n- material"),
        ],
        references=gen.build_references([
            {"file_name": "intro_to_heat.pdf", "subject": "heat", "page": 3},
            {"file_name": "intro_to_heat.pdf", "subject": "heat", "page": 7},
            {"file_name": "notes & more_100%.txt", "subject": "heat"},
        ]),
    )


class _AI:
    def __init__(self, text="", error=None, boom=False):
        self.text, self.error, self.boom, self.prompts = text, error, boom, []

    def generate(self, prompt, timeout=60):
        self.prompts.append(prompt)
        if self.boom:
            raise RuntimeError("down")
        return {"text": self.text, "error": self.error}


_GOOD_METHOD = """Here you go:
```json
{"research_question": "Does sleep affect exam scores?",
 "hypotheses": ["More sleep raises scores"],
 "design": "Observational cohort study",
 "independent_variables": ["hours of sleep"],
 "dependent_variables": "exam score",
 "controls": ["prior grades"],
 "data_collection": "Sleep diaries and exam records",
 "analysis": "Linear regression",
 "limitations": ["self-report bias"]}
```"""


# ---- format / references ---------------------------------------------------

@pytest.mark.parametrize("value,expected", [
    (None, "pdf"), ("", "pdf"), ("PDF", "pdf"), ("tex", "latex"), (".tex", "latex"),
    ("PowerPoint", "pptx"), ("slides", "pptx"), ("ppt", "pptx"),
])
def test_normalise_format(value, expected):
    assert gen.normalise_format(value) == expected


def test_unknown_format_is_a_clear_error():
    with pytest.raises(gen.GenerationError):
        gen.normalise_format("docx")


def test_references_merge_pages_per_file_and_have_unique_keys():
    refs = _model().references
    assert len(refs) == 2
    by_file = {r.file_name: r for r in refs}
    assert by_file["intro_to_heat.pdf"].pages == "3-7"
    assert len({r.key for r in refs}) == 2


# ---- PDF (#53) -------------------------------------------------------------

def test_pdf_opens_and_contains_title_sections_and_references(tmp_path):
    import pdfplumber
    path = gen.render_pdf(_model(), str(tmp_path / "r.pdf"))
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    assert "heat transfer & flow" in text
    assert "Overview" in text and "Conduction is 100% driven" in text
    assert "References" in text and "intro to heat" in text


# ---- LaTeX (#51) -----------------------------------------------------------

def test_latex_escapes_special_characters_in_body_and_bib(tmp_path):
    tex, bib = gen.render_latex(_model(), str(tmp_path / "r.tex"))
    body = open(tex, encoding="utf-8").read()
    assert r"100\% driven" in body and r"C\# and \$5" in body and r"\_under\_" in body
    assert r"heat transfer \& flow" in body
    assert r"\item temperature" in body
    assert "notes \\& more\\_100\\%" in open(bib, encoding="utf-8").read()
    assert bib.endswith(".bib") and os.path.exists(bib)


def test_latex_compiles_with_pdflatex(tmp_path):
    if not shutil.which("pdflatex"):
        pytest.skip("pdflatex not installed")
    tex, _ = gen.render_latex(_model(), str(tmp_path / "r.tex"))
    r = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "r.tex"],
                       cwd=str(tmp_path), capture_output=True, text=True, timeout=120)
    assert r.returncode == 0, r.stdout[-800:]
    assert (tmp_path / "r.pdf").exists()


def test_latex_with_bibtex_resolves_the_reference_list(tmp_path):
    if not (shutil.which("pdflatex") and shutil.which("bibtex")):
        pytest.skip("pdflatex/bibtex not installed")
    gen.render_latex(_model(), str(tmp_path / "r.tex"))
    run = lambda *c: subprocess.run(list(c), cwd=str(tmp_path), capture_output=True, text=True, timeout=120)
    assert run("pdflatex", "-interaction=nonstopmode", "r.tex").returncode == 0
    assert run("bibtex", "r").returncode == 0
    assert run("pdflatex", "-interaction=nonstopmode", "r.tex").returncode == 0
    assert "intro to heat" in (tmp_path / "r.bbl").read_text(encoding="utf-8")


def test_latex_without_references_has_no_bibliography_command(tmp_path):
    m = _model()
    m.references = []
    tex, bib = gen.render_latex(m, str(tmp_path / "r.tex"))
    assert "bibliography" not in open(tex, encoding="utf-8").read()
    assert bib == "" and not (tmp_path / "r.bib").exists()


# ---- PowerPoint (#52) ------------------------------------------------------

def _slides(path):
    from pptx import Presentation
    out = []
    for s in Presentation(path).slides:
        out.append((s.shapes.title.text, [sh.text_frame.text for sh in s.placeholders if sh.placeholder_format.idx == 1]))
    return out


def test_pptx_has_title_one_slide_per_section_and_sources(tmp_path):
    path = gen.render_pptx(_model(), str(tmp_path / "r.pptx"))
    slides = _slides(path)
    assert [t for t, _ in slides] == ["Literature review: heat transfer & flow", "Overview", "Variables", "Sources"]
    assert "Radiation needs no medium." in slides[1][1][0]
    assert "temperature" in slides[2][1][0]


def test_bullets_are_capped_in_number_and_length():
    long = " ".join(f"Sentence number {i} is here." for i in range(20))
    assert len(gen._bulletize(long)) == 5
    assert all(len(b) <= 181 for b in gen._bulletize("word " * 100 + "end."))


def test_llm_outline_is_used_when_valid():
    ai = _AI('{"slides": [{"title": "Big idea", "bullets": ["one", "two"]}, {"title": "", "bullets": ["x"]}]}')
    assert gen.outline_with_llm(ai, _model()) == [{"title": "Big idea", "bullets": ["one", "two"]}]


@pytest.mark.parametrize("ai", [_AI("not json at all"), _AI('{"slides": []}'), _AI(boom=True), _AI(error="x")])
def test_llm_outline_falls_back_to_deterministic_on_any_problem(ai):
    assert gen.outline_with_llm(ai, _model()) == gen.deterministic_outline(_model())


# ---- methodology (#66) -----------------------------------------------------

def test_methodology_json_is_extracted_from_fenced_chatter_and_validated():
    m = gen.generate_methodology(_AI(_GOOD_METHOD), "sleep and exam scores")
    assert m["research_question"].startswith("Does sleep")
    assert m["dependent_variables"] == ["exam score"]          # string coerced to list
    assert m["controls"] == ["prior grades"]


@pytest.mark.parametrize("text", [
    "I think you should study sleep.",
    "{not valid json}",
    '{"research_question": "q"}',                               # missing required fields
    '["a", "b"]',
    '{"research_question": "q", "design": "d", "independent_variables": [], "dependent_variables": ["y"]}',
])
def test_junk_model_output_raises_a_clear_message(text):
    with pytest.raises(gen.GenerationError) as e:
        gen.generate_methodology(_AI(text), "topic")
    assert "usable methodology" in str(e.value)


def test_model_down_is_a_different_message_from_junk():
    for ai in (_AI(boom=True), _AI(error="down"), _AI("")):
        with pytest.raises(gen.GenerationError) as e:
            gen.generate_methodology(ai, "topic")
        assert "language model" in str(e.value)


def test_methodology_needs_a_topic():
    with pytest.raises(gen.GenerationError):
        gen.generate_methodology(_AI(_GOOD_METHOD), "  ")


def test_methodology_renders_through_every_exporter(tmp_path):
    model = gen.methodology_to_report(gen.generate_methodology(_AI(_GOOD_METHOD), "sleep"))
    headings = [s.heading for s in model.sections]
    assert headings[:3] == ["Research question", "Hypotheses", "Study design"] and "Limitations" in headings
    for fmt, n in (("pdf", 1), ("latex", 1), ("pptx", 1)):      # a methodology has no references, so no .bib
        files = gen.render(model, fmt, str(tmp_path), "m")
        assert len(files) == n and all(os.path.getsize(f) > 0 for f in files)


# ---- request parsing -------------------------------------------------------

@pytest.mark.parametrize("raw,subject,fmt", [
    ("make a pdf report on heat transfer", "heat transfer", "pdf"),
    ("turn my machine learning notes into a powerpoint", "machine learning", "powerpoint"),
    ("generate a report about thermodynamics as a latex file", "thermodynamics", "latex"),
    ("make slides on neural nets", "neural nets", "slides"),
    ("make a report", "", None),
])
def test_parse_report_request(raw, subject, fmt):
    assert gen.parse_report_request(raw) == (subject, fmt)


def test_parse_methodology_request():
    assert gen.parse_methodology_request("draft a methodology for studying sleep and exam scores") == \
        "studying sleep and exam scores"


def test_slugify_is_filename_safe():
    assert gen.slugify("Heat & Mass / Transfer?!") == "heat-mass-transfer"
    assert gen.slugify("???") == "report"


def test_render_failure_leaves_a_clear_message(tmp_path, monkeypatch):
    monkeypatch.setattr(gen, "render_pdf", lambda *a, **k: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(gen.GenerationError) as e:
        gen.render(_model(), "pdf", str(tmp_path), "x")
    assert "nothing was saved" in str(e.value)


# ---- engine intents --------------------------------------------------------

class _FakeSynth:
    def __init__(self, review="Themes recur across the sources. They agree on basics.",
                 sources=("a.pdf", "b.pdf")):
        self.review, self.sources = review, list(sources)

    def generate_literature_review(self, subject):
        return {"review": self.review, "sources": self.sources}


@pytest.fixture
def engine():
    tmp = tempfile.mkdtemp()
    eng = make_engine(tmp)
    yield eng
    shutil.rmtree(tmp, ignore_errors=True)


def test_intents_are_declared_in_both_prefixed_and_stripped_forms(engine):
    for name in ("generate_report", "athena_generate_report", "methodology", "athena_methodology"):
        assert engine.can_handle(name)


@pytest.mark.parametrize("fmt,ext", [("pdf", ".pdf"), ("latex", ".tex"), ("pptx", ".pptx")])
def test_generate_report_creates_the_requested_file(engine, fmt, ext):
    engine.synthesis = _FakeSynth()
    r = engine.handle("athena_generate_report", {"subject": "heat", "format": fmt}, {})
    assert r["confidence"] > 0.5 and r["data"]["format"] == fmt
    assert any(f.endswith(ext) for f in r["data"]["files"])
    assert all(os.path.exists(f) for f in r["data"]["files"])
    assert os.path.basename(os.path.dirname(r["data"]["files"][0])) == "exports"


def test_generate_report_reads_subject_and_format_from_the_raw_query(engine):
    engine.synthesis = _FakeSynth()
    r = engine.handle("generate_report", {}, {"raw_query": "make a powerpoint on heat transfer"})
    assert r["data"]["format"] == "pptx" and r["data"]["subject"] == "heat transfer"


def test_generate_report_asks_when_no_subject(engine):
    r = engine.handle("generate_report", {}, {"raw_query": "make a report"})
    assert "Which subject" in r["response"] and r["confidence"] == 0.0


def test_generate_report_with_nothing_indexed_creates_no_file(engine):
    engine.synthesis = _FakeSynth(review="I don't have any documents ingested under \"x\" to review.", sources=[])
    r = engine.handle("generate_report", {"subject": "x"}, {})
    assert "nothing to report on" in r["response"] and "files" not in r["data"]
    assert not os.path.exists(engine._export_dir())


def test_generate_report_when_the_model_failed_creates_no_file(engine):
    engine.synthesis = _FakeSynth(review="I had trouble generating a review for \"x\".")
    r = engine.handle("generate_report", {"subject": "x"}, {})
    assert "didn't create a file" in r["response"] and not os.path.exists(engine._export_dir())


def test_generate_report_rejects_unknown_format_politely(engine):
    engine.synthesis = _FakeSynth()
    r = engine.handle("generate_report", {"subject": "x", "format": "docx"}, {})
    assert "PDF, a LaTeX file or a PowerPoint" in r["response"]


def test_methodology_intent_writes_a_file_from_a_fake_model(engine):
    engine.llm = _AI(_GOOD_METHOD)
    r = engine.handle("athena_methodology", {"question": "sleep and exam scores", "format": "pdf"}, {})
    assert r["data"]["methodology"]["design"] == "Observational cohort study"
    assert os.path.exists(r["data"]["files"][0])


def test_methodology_intent_with_junk_model_output_says_so_and_saves_nothing(engine):
    engine.llm = _AI("lol no")
    r = engine.handle("methodology", {"question": "sleep"}, {})
    assert "usable methodology" in r["response"] and not os.path.exists(engine._export_dir())


def test_methodology_topic_comes_from_raw_query_when_entities_are_empty(engine):
    ai = _AI(_GOOD_METHOD)
    engine.llm = ai
    engine.handle("methodology", {}, {"raw_query": "draft a methodology for studying sleep"})
    assert "studying sleep" in ai.prompts[0]
